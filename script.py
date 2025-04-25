import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from io import BytesIO

# 设置项目根目录
project_root = r'.'
output_file = 'ProjectSourceCode.docx'
max_size_kb = 40  # KB

# 想要包含的代码后缀
valid_extensions = ('.java', '.js', '.jsx', '.ts', '.tsx', '.json', '.html', '.css')

# 初始化 Word 文档
doc = Document()
doc.styles['Normal'].font.name = 'Courier New'
doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
doc.styles['Normal'].font.size = Pt(9)

def get_doc_size_kb(document):
    tmp_stream = BytesIO()
    document.save(tmp_stream)
    size_kb = len(tmp_stream.getvalue()) / 1024
    return size_kb

stop = False
for root, dirs, files in os.walk(project_root):
    for file in files:
        if file.endswith(valid_extensions):
            file_path = os.path.join(root, file)
            relative_path = os.path.relpath(file_path, project_root)
            doc.add_heading(relative_path, level=2)
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                doc.add_paragraph(content)
            except Exception as e:
                doc.add_paragraph(f'[Error reading file: {e}]')

            # 检查当前文档大小
            if get_doc_size_kb(doc) > max_size_kb:
                print(f"⚠️ 文档已超过 {max_size_kb}KB，导出中止。")
                stop = True
                break
    if stop:
        break

# 保存 Word 文档
doc.save(output_file)
print(f"✅ 源代码已导出到：{output_file}")